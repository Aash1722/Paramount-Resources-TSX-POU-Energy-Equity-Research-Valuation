"""
Paramount Resources (TSX:POU) - Module 4: Sum-of-the-Parts Net Asset Value
Willesden Green off the Module 3 type curve; Sinclair as a risked development
project; Kaybob as a declining annuity; land at a per-acre mark.
Every asset value is a closed-form annuity so the sensitivity grid stays live.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
F = "Arial"
BLUE  = Font(name=F, size=10, color="0000FF")
BLACK = Font(name=F, size=10)
GREEN = Font(name=F, size=10, color="008000")
BOLD  = Font(name=F, size=10, bold=True)
BOLDW = Font(name=F, size=10, bold=True, color="FFFFFF")
RED   = Font(name=F, size=10, italic=True, color="C00000")
TITLE = Font(name=F, size=14, bold=True, color="1F3864")
SUB   = Font(name=F, size=10, italic=True, color="595959")
HDR   = PatternFill("solid", fgColor="1F3864")
SECF  = PatternFill("solid", fgColor="D9E2F3")
YEL   = PatternFill("solid", fgColor="FFFF00")
GRN   = PatternFill("solid", fgColor="E2EFDA")
TOP   = Border(top=Side(style="thin", color="404040"))
CUR   = '$#,##0;($#,##0);-'
CUR1  = '$#,##0.00;($#,##0.00);-'
CUR2  = '$#,##0.0;($#,##0.0);-'
NUM   = '#,##0;(#,##0);-'
NUM1  = '#,##0.0;(#,##0.0);-'
PCT   = '0.0%'
PCT0  = '0%'
MULT  = '0.00"x"'

R = {}
def title(ws, n, s):
    ws["A1"] = n; ws["A1"].font = TITLE
    ws["A2"] = s; ws["A2"].font = SUB
    ws.freeze_panes = "A4"
def sec(ws, r, t, w=7):
    ws.cell(r, 1, t).font = Font(name=F, size=10, bold=True, color="1F3864")
    for c in range(1, w+1): ws.cell(r, c).fill = SECF
    return r+1
def row(ws, r, key, label, val=None, formula=None, fmt=NUM1, note=None,
        total=False, fill=None, font=None, col=2):
    c = ws.cell(r, 1, label); c.font = font or (BOLD if total else BLACK)
    cell = ws.cell(r, col, formula if formula is not None else val)
    cell.font = BLACK if formula is not None else BLUE
    if total: cell.font = BOLD; cell.border = TOP
    cell.number_format = fmt
    if fill: cell.fill = fill
    if note:
        n = ws.cell(r, 5, note); n.font = SUB
        n.alignment = Alignment(wrap_text=True, vertical="top")
        if len(note) > 75: ws.row_dimensions[r].height = 28
    R[key] = r
    return r+1
def W(ws, s):
    for k, v in s.items(): ws.column_dimensions[k].width = v

A = "Assumptions!"
def a(key): return f"{A}$B${R[key]}"

# =====================================================================
# README
# =====================================================================
ws = wb.active; ws.title = "README"
title(ws, "Paramount Resources — Module 4: Sum-of-the-Parts Net Asset Value",
      "Values each asset separately, then rolls to a per-share number and compares it to the market price.")
W(ws, {"A": 30, "B": 104})
r = 4
r = sec(ws, r, "APPROACH", 2)
for k, v in [
    ("Why sum-of-the-parts", "Paramount's assets have completely different economics: Willesden Green is a producing liquids-rich play, Sinclair is an unbuilt dry-gas project, Kaybob is a declining annuity, and the land bank generates nothing. A single corporate multiple would obscure all of that."),
    ("Willesden Green", "Producing base as a declining annuity, plus the present value of a 20-year drilling programme valued off the Module 3 type curve. Net present value per well is live to the price deck rather than hardcoded."),
    ("Sinclair", "Treated as a development project: remaining construction capital deducted, then a delayed annuity from first gas, multiplied by an execution risk factor. This is the asset that decides the valuation."),
    ("Kaybob", "Declining annuity net of the maintenance capital required to sustain it."),
    ("Land", "Acreage carrying no reserves, at a per-acre mark. The softest number in the model — treat it as option value, not value."),
    ("Closed-form maths", "Every asset uses annuity algebra rather than a year-by-year grid. That keeps the sensitivity table on NAV_Summary fully live instead of static."),
]:
    ws.cell(r, 1, k).font = BOLD
    c = ws.cell(r, 2, v); c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 44 if len(v) > 150 else 30
    r += 1
r += 1
r = sec(ws, r, "THE TWO FORMULAS DOING ALL THE WORK", 2)
for k, v in [
    ("Declining annuity", "PV = Q x 365 x netback / (r + d). A stream declining at rate d, discounted at r. This is why the decline rate matters as much as the discount rate — both sit in the denominator."),
    ("Level annuity", "PV = C x [1 - (1+r)^-n] / r. Used for the drilling programme: a fixed number of wells drilled each year for n years."),
]:
    ws.cell(r, 1, k).font = BOLD
    c = ws.cell(r, 2, v); c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 30; r += 1
r += 1
r = sec(ws, r, "HEALTH WARNINGS", 2)
for v in [
    "This is a pre-tax-pool NAV with a single effective tax drag applied at the corporate level. A full model would schedule the tax pools remaining from the Grande Prairie disposition explicitly.",
    "The Sinclair execution risk factor is a judgement, not an observation. It is the single largest lever in the file — vary it before quoting any conclusion.",
    "Land value is the weakest input. Frontier acreage in Horn River, Liard and the Northwest Territories may be worth nothing at current prices.",
    "Module 3 showed the drilling programme assumes no well-to-well interference, so the well count is a floor and the programme value is therefore optimistic at the margin.",
]:
    c = ws.cell(r, 1, "•  " + v); c.font = SUB
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 28; r += 1
r += 1
ws.cell(r, 1, "Hypothetical case study. Not investment advice.").font = SUB

# =====================================================================
# ASSUMPTIONS
# =====================================================================
ws = wb.create_sheet("Assumptions")
title(ws, "Valuation Assumptions",
      "Yellow cells are the levers. Netbacks are computed from the price deck and each asset's product mix.")
W(ws, {"A": 46, "B": 14, "C": 3, "D": 2, "E": 58})
r = 4
r = sec(ws, r, "PRICE DECK AND DISCOUNTING")
r = row(ws, r, "wti", "WTI (US$/Bbl)", 80.0, fmt=CUR1, fill=YEL,
        note="Module 2 Base deck long-run level. Spot is near US$86 on a geopolitical premium.")
r = row(ws, r, "fx", "CAD/USD", 1.36, fmt='0.000', fill=YEL)
r = row(ws, r, "gas", "Realised gas (C$/Mcf)", 4.33, fmt=CUR1, fill=YEL,
        note="AECO C$3.20/GJ plus C$0.90 diversification premium, converted at 1.055.")
r = row(ws, r, "ngl_r", "Other NGLs as % of condensate", 0.32, fmt=PCT0, fill=YEL)
r = row(ws, r, "disc", "Discount rate", 0.10, fmt=PCT, fill=YEL,
        note="Industry convention for reserve valuation is 10% before tax.")
r = row(ws, r, "shares", "Shares outstanding (mm)", 145.8, fmt=NUM1)
r = row(ws, r, "price", "Current share price (C$)", 32.65, fmt=CUR1, fill=YEL)
r += 1

r = sec(ws, r, "ASSET MIX AND UNIT COSTS")
hdr = r
ws.cell(r, 1, "").font = BOLDW
for i, h in enumerate(["Willesden Grn", "Sinclair", "Kaybob"]):
    c = ws.cell(r, 2+i, h); c.font = BOLDW; c.alignment = Alignment(horizontal="right")
for c in range(1, 5): ws.cell(r, c).fill = HDR
r += 1
def arow(r, key, label, vals, fmt, note=None):
    ws.cell(r, 1, label).font = BLACK
    for i, v in enumerate(vals):
        c = ws.cell(r, 2+i, v); c.font = BLUE; c.number_format = fmt; c.fill = YEL
    if note:
        n = ws.cell(r, 5, note); n.font = SUB
        n.alignment = Alignment(wrap_text=True, vertical="top")
        if len(note) > 75: ws.row_dimensions[r].height = 28
    R[key] = r
    return r+1
r = arow(r, "liq", "Liquids % of production", [0.55, 0.10, 0.40], PCT0,
         "Sinclair is dry gas — this single row explains most of the valuation gap between the assets.")
r = arow(r, "cnd", "Condensate share of liquids", [0.80, 0.75, 0.70], PCT0)
r = arow(r, "opex", "Operating expense (C$/Boe)", [4.71, 4.00, 11.00], CUR1,
         "Willesden Green is the Q4 2025 actual. Sinclair assumes high-rate dry gas through an owned plant.")
r = arow(r, "tran", "Transportation (C$/Boe)", [4.75, 4.50, 5.00], CUR1)
r = arow(r, "roy", "Royalty rate", [0.06, 0.06, 0.08], PCT0)
r += 1

r = sec(ws, r, "COMPUTED NETBACK BY ASSET (C$/Boe)")
nb_hdr = r
ws.cell(r, 1, "").font = BOLDW
for i, h in enumerate(["Willesden Grn", "Sinclair", "Kaybob"]):
    c = ws.cell(r, 2+i, h); c.font = BOLDW; c.alignment = Alignment(horizontal="right")
for c in range(1, 5): ws.cell(r, c).fill = HDR
r += 1
def crow(r, key, label, f, fmt, note=None, total=False):
    ws.cell(r, 1, label).font = BOLD if total else BLACK
    for i in range(3):
        col = get_column_letter(2+i)
        c = ws.cell(r, 2+i, f(col)); c.font = BLACK; c.number_format = fmt
        if total: c.font = BOLD; c.border = TOP; c.fill = GRN
    if note:
        n = ws.cell(r, 5, note); n.font = SUB
        n.alignment = Alignment(wrap_text=True, vertical="top")
        if len(note) > 75: ws.row_dimensions[r].height = 28
    R[key] = r
    return r+1
r = crow(r, "cond_p", "Condensate price (C$/Bbl)",
         lambda c: f"={a('wti')}*{a('fx')}", CUR1)
r = crow(r, "rev", "Revenue per Boe",
         lambda c: (f"={c}{R['liq']}*{c}{R['cnd']}*{c}{R['cond_p']}"
                    f"+{c}{R['liq']}*(1-{c}{R['cnd']})*{c}{R['cond_p']}*{a('ngl_r')}"
                    f"+(1-{c}{R['liq']})*6*{a('gas')}"), CUR1,
         "One Boe = liquids fraction in barrels, plus the gas remainder times six Mcf.")
r = crow(r, "nb", "Netback per Boe",
         lambda c: (f"={c}{R['rev']}*(1-{c}{R['roy']})-{c}{R['opex']}-{c}{R['tran']}"), CUR1,
         "Drives every asset value below.", total=True)
r += 1

r = sec(ws, r, "WILLESDEN GREEN — PRODUCING BASE AND DRILLING PROGRAMME")
r = row(ws, r, "wg_q", "Current production (Boe/d)", 27466, fmt=NUM,
        note="Q2 2026 actual.")
r = row(ws, r, "wg_d", "Base decline rate", 0.28, fmt=PCT, fill=YEL,
        note="Corporate-level decline on the producing base, gentler than a single new well.")
r = row(ws, r, "wg_eur", "EUR per well (MBoe)", 1441, fmt=NUM,
        note="MODULE 3 OUTPUT — calibrated to the disclosed 210-day rate.")
r = row(ws, r, "wg_pvf", "Production PV factor", 0.57, fmt='0.00', fill=YEL,
        note="Collapses the timing of a well's production into one discount factor. Calibrated so NPV per well reproduces the Module 3 full-cycle figure of C$23.5mm.")
r = row(ws, r, "wg_cost", "Full-cycle cost per well (C$mm)", 18.0, fmt=CUR2, fill=YEL,
        note="MODULE 3: C$13mm drilling and completion plus C$5mm allocated facility capital.")
r = row(ws, r, "wg_npv", "Net present value per well (C$mm)",
        formula=f"={a('wg_eur')}*1000*B{R['nb']}*{a('wg_pvf')}/1000000-{a('wg_cost')}",
        fmt=CUR2, note="Live to price. Cross-check against Module 3's C$23.5mm at WTI US$80.", fill=GRN)
r = row(ws, r, "wg_n", "Wells drilled per year", 16, fmt=NUM,
        note="MODULE 3: the count that sustains the 50,000 Boe/d plateau. Assumes no well interference, so it is a floor.")
r = row(ws, r, "wg_yrs", "Programme life (years)", 20, fmt=NUM, fill=YEL,
        note="Company states the plateau is sustainable for 20+ years.")
r = row(ws, r, "wg_risk", "Risk factor on undrilled locations", 0.80, fmt=PCT0, fill=YEL,
        note="Haircut for locations not yet booked as proved developed.")
r += 1

r = sec(ws, r, "SINCLAIR — DEVELOPMENT PROJECT")
r = row(ws, r, "sc_q", "Plateau production (Boe/d)", 55000, fmt=NUM,
        note="Constrained by 335 MMcf/d of firm egress, not by the resource.")
r = row(ws, r, "sc_plat", "Years held at plateau", 10, fmt=NUM, fill=YEL,
        note="Egress is capped at 335 MMcf/d while the resource is far larger, so the asset is designed to hold plateau rather than decline from day one. Modelling immediate decline understates it badly.")
r = row(ws, r, "sc_d", "Decline rate after plateau", 0.15, fmt=PCT, fill=YEL)
r = row(ws, r, "sc_capex", "Remaining construction capital (C$mm)", 800, fmt=CUR, fill=YEL,
        note="Roughly C$360mm in 2026 plus C$440mm in 2027 per company guidance.")
r = row(ws, r, "sc_start", "Years to first gas", 1.4, fmt=NUM1, fill=YEL,
        note="Q4 2027 from an August 2026 valuation date.")
r = row(ws, r, "sc_ramp", "Years to reach plateau", 2.0, fmt=NUM1, fill=YEL)
r = row(ws, r, "sc_maint", "Maintenance capital (% of netback)", 0.30, fmt=PCT0, fill=YEL,
        note="Dry gas needs continuous drilling to hold plateau.")
r = row(ws, r, "sc_risk", "Execution risk factor", 0.70, fmt=PCT0, fill=YEL,
        note="THE BIGGEST LEVER IN THE MODEL. Single asset, single plant, unhedged gas price at start-up. Vary this from 50% to 90% before quoting anything.")
r += 1

r = sec(ws, r, "KAYBOB AND OTHER")
r = row(ws, r, "kb_q", "Kaybob production (Boe/d)", 19413, fmt=NUM, note="Q2 2026 actual.")
r = row(ws, r, "kb_d", "Kaybob decline rate", 0.12, fmt=PCT, fill=YEL)
r = row(ws, r, "kb_maint", "Kaybob maintenance capital (% of netback)", 0.45, fmt=PCT0, fill=YEL,
        note="Higher than Willesden Green: mature wells, worse economics.")
r = row(ws, r, "land_ac", "Undeveloped net acreage (thousand acres)", 2349, fmt=NUM,
        note="Acreage carrying no reserves at 31-Dec-2025.")
r = row(ws, r, "land_val", "Land value (C$ per acre)", 150, fmt=CUR, fill=YEL,
        note="WEAKEST INPUT IN THE FILE. Frontier acreage may be worth nothing. Sensitivity runs C$0 to C$400.")
r += 1

r = sec(ws, r, "CORPORATE")
r = row(ws, r, "ga", "General and administrative (C$mm per year)", 60, fmt=CUR, fill=YEL)
r = row(ws, r, "cash", "Net cash at 30-Jun-2026 (C$mm)", 329.1, fmt=CUR,
        note="Module 1. Net cash, so it adds to value.")
r = row(ws, r, "inv", "Investments in securities (C$mm)", 208.0, fmt=CUR,
        note="Confirm the residual after the AKITA distribution on 16-Jul-2026.")
r = row(ws, r, "aro", "Present value of abandonment beyond reserves (C$mm)", 300, fmt=CUR, fill=YEL,
        note="Long-dated obligations on legacy properties not captured in the reserve report.")
r = row(ws, r, "tax", "Effective tax drag on asset value", 0.12, fmt=PCT, fill=YEL,
        note="Low because large tax pools remain from the Grande Prairie disposition. A full model would schedule the pools explicitly.")

# =====================================================================
# NAV BUILD
# =====================================================================
ws = wb.create_sheet("NAV_Build")
title(ws, "Asset-by-Asset Net Asset Value",
      "Closed-form annuities. Every value is live to the Assumptions sheet.")
W(ws, {"A": 50, "B": 15, "C": 3, "D": 2, "E": 56})
r = 4
r = sec(ws, r, "WILLESDEN GREEN")
r = row(ws, r, "n_wg1", "Producing base — declining annuity (C$mm)",
        formula=f"={a('wg_q')}*365*Assumptions!B{R['nb']}/({a('disc')}+{a('wg_d')})/1000000",
        fmt=CUR, note="Q x 365 x netback / (r + d).")
r = row(ws, r, "n_wg2", "Drilling programme — level annuity (C$mm)",
        formula=(f"={a('wg_n')}*{a('wg_npv')}*{a('wg_risk')}"
                 f"*(1-(1+{a('disc')})^-{a('wg_yrs')})/{a('disc')}"),
        fmt=CUR, note="Wells per year x risked net present value per well, as a 20-year annuity.")
r = row(ws, r, "n_wg", "Willesden Green total (C$mm)",
        formula=f"=B{R['n_wg1']}+B{R['n_wg2']}", fmt=CUR, total=True, fill=GRN)
r += 1

r = sec(ws, r, "SINCLAIR")
r = row(ws, r, "n_sc0", "Annual cash flow at plateau (C$mm)",
        formula=(f"={a('sc_q')}*365*Assumptions!C{R['nb']}*(1-{a('sc_maint')})/1000000"),
        fmt=CUR, note="Net of the maintenance capital needed to hold the plateau.")
r = row(ws, r, "n_sc1", "Value at first gas: plateau annuity plus declining tail (C$mm)",
        formula=(f"=B{R['n_sc0']}*(1-(1+{a('disc')})^-{a('sc_plat')})/{a('disc')}"
                 f"+B{R['n_sc0']}/({a('disc')}+{a('sc_d')})*(1+{a('disc')})^-{a('sc_plat')}"),
        fmt=CUR, note="Level annuity across the plateau years, then a declining perpetuity discounted back.")
r = row(ws, r, "n_sc2", "Discounted to today (C$mm)",
        formula=f"=B{R['n_sc1']}/(1+{a('disc')})^({a('sc_start')}+{a('sc_ramp')}/2)",
        fmt=CUR, note="Start-up delay plus half the ramp period.")
r = row(ws, r, "n_sc3", "Less: remaining construction capital (C$mm)",
        formula=f"=-{a('sc_capex')}/(1+{a('disc')})^({a('sc_start')}/2)", fmt=CUR)
r = row(ws, r, "n_sc", "Sinclair total, risked (C$mm)",
        formula=f"=(B{R['n_sc2']}+B{R['n_sc3']})*{a('sc_risk')}", fmt=CUR, total=True, fill=GRN,
        note="Risk factor applied to the net project value, not to the gross cash flow.")
r += 1

r = sec(ws, r, "KAYBOB, LAND AND CORPORATE")
r = row(ws, r, "n_kb", "Kaybob — declining annuity net of maintenance (C$mm)",
        formula=(f"={a('kb_q')}*365*Assumptions!D{R['nb']}*(1-{a('kb_maint')})"
                 f"/({a('disc')}+{a('kb_d')})/1000000"), fmt=CUR, total=True, fill=GRN)
r = row(ws, r, "n_land", "Undeveloped land (C$mm)",
        formula=f"={a('land_ac')}*1000*{a('land_val')}/1000000", fmt=CUR, total=True, fill=GRN)
r = row(ws, r, "n_ga", "Corporate overhead — perpetuity (C$mm)",
        formula=f"=-{a('ga')}/{a('disc')}", fmt=CUR,
        note="Capitalised head office cost. Frequently omitted, which flatters NAV.")
r = row(ws, r, "n_aro", "Abandonment beyond reserves (C$mm)",
        formula=f"=-{a('aro')}", fmt=CUR)
r += 1

r = sec(ws, r, "ROLL-UP")
r = row(ws, r, "n_gross", "Gross asset value before tax (C$mm)",
        formula=(f"=B{R['n_wg']}+B{R['n_sc']}+B{R['n_kb']}+B{R['n_land']}"
                 f"+B{R['n_ga']}+B{R['n_aro']}"), fmt=CUR, total=True)
r = row(ws, r, "n_tax", "Less: tax drag (C$mm)",
        formula=f"=-MAX(0,B{R['n_gross']})*{a('tax')}", fmt=CUR)
r = row(ws, r, "n_net", "After-tax asset value (C$mm)",
        formula=f"=B{R['n_gross']}+B{R['n_tax']}", fmt=CUR, total=True)
r = row(ws, r, "n_cash", "Plus: net cash and investments (C$mm)",
        formula=f"={a('cash')}+{a('inv')}", fmt=CUR)
r = row(ws, r, "n_nav", "NET ASSET VALUE (C$mm)",
        formula=f"=B{R['n_net']}+B{R['n_cash']}", fmt=CUR, total=True, fill=YEL)
r = row(ws, r, "n_nps", "NAV per share (C$)",
        formula=f"=B{R['n_nav']}/{a('shares')}", fmt=CUR1, total=True, fill=YEL)
r = row(ws, r, "n_pnav", "Price / NAV",
        formula=f"={a('price')}/B{R['n_nps']}", fmt=MULT, total=True, fill=YEL,
        note="Above 1.0x means the market pays more than this model supports.")
r = row(ws, r, "n_up", "Upside / (downside) to NAV",
        formula=f"=B{R['n_nps']}/{a('price')}-1", fmt=PCT, total=True, fill=YEL)

NB = R  # snapshot

# =====================================================================
# NAV SUMMARY
# =====================================================================
ws = wb.create_sheet("NAV_Summary")
title(ws, "Valuation Summary",
      "Sum of the parts, contribution by asset, and live sensitivity to the two assumptions that matter most.")
W(ws, {"A": 42, "B": 14, "C": 12, "D": 12, "E": 12, "F": 12, "G": 12, "H": 3, "I": 46})
r = 4
r = sec(ws, r, "SUM OF THE PARTS", 7)
ws.cell(r, 1, "Asset").font = BOLDW
for i, h in enumerate(["C$mm", "C$/share", "% of gross"]):
    c = ws.cell(r, 2+i, h); c.font = BOLDW; c.alignment = Alignment(horizontal="right")
for c in range(1, 8): ws.cell(r, c).fill = HDR
r += 1
parts_start = r
for lab, key in [("Willesden Green", "n_wg"), ("Sinclair (risked)", "n_sc"),
                 ("Kaybob", "n_kb"), ("Undeveloped land", "n_land"),
                 ("Corporate overhead", "n_ga"), ("Abandonment", "n_aro")]:
    ws.cell(r, 1, lab).font = BLACK
    c = ws.cell(r, 2, f"=NAV_Build!B{NB[key]}"); c.font = GREEN; c.number_format = CUR
    c = ws.cell(r, 3, f"=B{r}/{a('shares')}"); c.font = BLACK; c.number_format = CUR1
    c = ws.cell(r, 4, f"=B{r}/SUM($B${parts_start}:$B${parts_start+5})"); c.font = BLACK; c.number_format = PCT0
    r += 1
ws.cell(r, 1, "Gross asset value").font = BOLD
for cc, fmt in [(2, CUR), (3, CUR1)]:
    c = ws.cell(r, cc, f"=SUM({get_column_letter(cc)}{parts_start}:{get_column_letter(cc)}{r-1})")
    c.font = BOLD; c.number_format = fmt; c.border = TOP
r += 1
for lab, key in [("Tax drag", "n_tax"), ("Net cash and investments", "n_cash")]:
    ws.cell(r, 1, lab).font = BLACK
    c = ws.cell(r, 2, f"=NAV_Build!B{NB[key]}"); c.font = GREEN; c.number_format = CUR
    c = ws.cell(r, 3, f"=B{r}/{a('shares')}"); c.font = BLACK; c.number_format = CUR1
    r += 1
ws.cell(r, 1, "NET ASSET VALUE").font = BOLD
for cc, fmt in [(2, CUR), (3, CUR1)]:
    c = ws.cell(r, cc, f"=NAV_Build!{'B' if cc==2 else 'B'}{NB['n_nav'] if cc==2 else NB['n_nps']}")
    c.font = BOLD; c.number_format = fmt; c.border = TOP; c.fill = YEL
nav_row = r
r += 2

r = sec(ws, r, "VERSUS THE MARKET", 7)
for lab, f_, fmt, note in [
    ("NAV per share (C$)", f"=NAV_Build!B{NB['n_nps']}", CUR1, None),
    ("Current share price (C$)", f"={a('price')}", CUR1, None),
    ("Price / NAV", f"=NAV_Build!B{NB['n_pnav']}", MULT,
     "Compare with 1.46x against before-tax 2P NPV10 in the Data Book. This NAV adds land and Sinclair upside but deducts tax, overhead and abandonment."),
    ("Upside / (downside)", f"=NAV_Build!B{NB['n_up']}", PCT, None),
]:
    ws.cell(r, 1, lab).font = BOLD
    c = ws.cell(r, 2, f_); c.font = GREEN; c.number_format = fmt; c.fill = YEL
    if note:
        n = ws.cell(r, 9, note); n.font = SUB
        n.alignment = Alignment(wrap_text=True, vertical="top"); ws.row_dimensions[r].height = 40
    r += 1
r += 1

# ---- live sensitivity grid: WTI x Sinclair risk factor ----
r = sec(ws, r, "SENSITIVITY — NAV PER SHARE (C$)", 7)
ws.cell(r, 1, "WTI (US$/Bbl)  \\  Sinclair execution risk").font = BOLD
grid_hdr = r
risks = [0.5, 0.6, 0.7, 0.8, 0.9]
for i, rk in enumerate(risks):
    c = ws.cell(r, 2+i, rk); c.font = BOLDW; c.number_format = PCT0
    c.alignment = Alignment(horizontal="right"); c.fill = HDR
ws.cell(r, 1).fill = HDR; ws.cell(r, 1).font = BOLDW
r += 1
wtis = [55, 65, 75, 85, 95]

def nav_formula(wti_ref, risk_ref):
    """Rebuild the whole NAV inline so the grid is live."""
    d, fx, gas, nglr = a('disc'), a('fx'), a('gas'), a('ngl_r')
    def nb(col, liq, cnd, opx, trn, roy):
        cond = f"({wti_ref}*{fx})"
        rev = (f"({liq}*{cnd}*{cond}+{liq}*(1-{cnd})*{cond}*{nglr}+(1-{liq})*6*{gas})")
        return f"({rev}*(1-{roy})-{opx}-{trn})"
    L, C, O, T, Y = (f"Assumptions!$B${R['liq']}", f"Assumptions!$B${R['cnd']}",
                     f"Assumptions!$B${R['opex']}", f"Assumptions!$B${R['tran']}",
                     f"Assumptions!$B${R['roy']}")
    nb_wg = nb("B", L, C, O, T, Y)
    nb_sc = nb("C", f"Assumptions!$C${R['liq']}", f"Assumptions!$C${R['cnd']}",
               f"Assumptions!$C${R['opex']}", f"Assumptions!$C${R['tran']}", f"Assumptions!$C${R['roy']}")
    nb_kb = nb("D", f"Assumptions!$D${R['liq']}", f"Assumptions!$D${R['cnd']}",
               f"Assumptions!$D${R['opex']}", f"Assumptions!$D${R['tran']}", f"Assumptions!$D${R['roy']}")
    wg1 = f"{a('wg_q')}*365*{nb_wg}/({d}+{a('wg_d')})/1000000"
    npvw = f"({a('wg_eur')}*1000*{nb_wg}*{a('wg_pvf')}/1000000-{a('wg_cost')})"
    wg2 = f"{a('wg_n')}*{npvw}*{a('wg_risk')}*(1-(1+{d})^-{a('wg_yrs')})/{d}"
    cf = f"({a('sc_q')}*365*{nb_sc}*(1-{a('sc_maint')})/1000000)"
    scv = (f"({cf}*(1-(1+{d})^-{a('sc_plat')})/{d}"
           f"+{cf}/({d}+{a('sc_d')})*(1+{d})^-{a('sc_plat')})")
    sc1 = (f"({scv}/(1+{d})^({a('sc_start')}+{a('sc_ramp')}/2)"
           f"-{a('sc_capex')}/(1+{d})^({a('sc_start')}/2))*{risk_ref}")
    kb = f"{a('kb_q')}*365*{nb_kb}*(1-{a('kb_maint')})/({d}+{a('kb_d')})/1000000"
    land = f"{a('land_ac')}*1000*{a('land_val')}/1000000"
    gross = f"({wg1}+{wg2}+{sc1}+{kb}+{land}-{a('ga')}/{d}-{a('aro')})"
    return (f"=({gross}*(1-{a('tax')})+{a('cash')}+{a('inv')})/{a('shares')}")

grid_start = r
for i, w in enumerate(wtis):
    c = ws.cell(r, 1, w); c.font = BOLDW; c.number_format = CUR; c.fill = HDR
    c.alignment = Alignment(horizontal="right")
    for j in range(len(risks)):
        rk_ref = f"{get_column_letter(2+j)}${grid_hdr}"
        wt_ref = f"$A{r}"
        cell = ws.cell(r, 2+j, nav_formula(wt_ref, rk_ref))
        cell.font = BLACK; cell.number_format = CUR1
    r += 1
r += 1
c = ws.cell(r, 1, "Cells shaded nothing: read against the current share price above. "
                  "The grid is live — change any Assumptions input and it moves.")
c.font = SUB
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
r += 2

r = sec(ws, r, "HOW TO READ THIS", 7)
for t in [
    "1. Look at the share of gross value coming from Sinclair. If it is large, this is not a producing-company valuation — it is a bet on one project being delivered on time.",
    "2. Compare NAV per share to the C$32.65 market price. A price above NAV means the market is more optimistic than these assumptions.",
    "3. Find the combination of WTI and execution risk that justifies today's price. State it explicitly. That combination IS the market's implied view, and agreeing or disagreeing with it is the investment thesis.",
]:
    c = ws.cell(r, 1, t); c.font = BLACK
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 30; r += 1

for s in wb.worksheets:
    s.sheet_view.showGridLines = False
wb.save("/home/claude/pou/POU_NAV_v1.xlsx")
print("saved")
